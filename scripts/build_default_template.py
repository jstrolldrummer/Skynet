"""Generates templates/proposal_template.docx — the Wyatt + Gray default.

Idempotent: re-run this whenever you want to refresh the bundled template.
The generated .docx contains every reusable boilerplate section from the
sample W&G proposal plus {{PLACEHOLDERS}} for the per-project bits the
bid-watcher fills in.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "templates" / "proposal_template.docx"


TEAL = RGBColor(0x0E, 0x6E, 0x82)


def _set_run(run, *, bold=False, size=None, color=None):
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run(run, bold=True, size=12, color=TEAL)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def _item_heading(doc, code_and_title, total="$ TBD"):
    p = doc.add_paragraph()
    run = p.add_run(code_and_title)
    _set_run(run, bold=True, size=11)
    tab = p.add_run("\t" + total)
    _set_run(tab, bold=True, size=11)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def _para(doc, text, bold=False, size=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run(run, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def _bullets(doc, lines):
    for line in lines:
        p = doc.add_paragraph(style="List Bullet")
        run = p.runs[0] if p.runs else p.add_run("")
        run.text = line
        _set_run(run, size=10)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # ---------- COVER PAGE ----------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("{{PROJECT_TITLE}}")
    _set_run(run, bold=True, size=28)

    sub = doc.add_paragraph()
    sr = sub.add_run("{{PROPOSAL_NUMBER}}")
    _set_run(sr, size=14, color=TEAL)

    doc.add_paragraph()  # spacer

    # Prepared By / Prepared For — two-column table
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells

    def _fill_party(cell, header_text, lines):
        p = cell.paragraphs[0]
        r = p.add_run(header_text)
        _set_run(r, bold=True, size=11)
        for line in lines:
            cell.add_paragraph(line)

    _fill_party(left, "Prepared By", [
        "{{SENDER_NAME}}",
        "{{SENDER_COMPANY}}",
        "{{SENDER_PHONE}}",
        "{{SENDER_EMAIL}}",
        "{{SENDER_ADDRESS}}",
    ])
    _fill_party(right, "Prepared For", [
        "{{CONTACT_NAME}}",
        "{{CLIENT_COMPANY}}",
        "{{CONTACT_EMAIL}}",
        "{{PROJECT_LOCATION}}",
    ])

    doc.add_page_break()

    # ---------- HEADER ROW ----------
    header_p = doc.add_paragraph()
    r1 = header_p.add_run("{{PROPOSAL_NUMBER}}")
    _set_run(r1, bold=True, size=14)
    header_p.add_run("\nIssue Date {{ISSUE_DATE}}\n")
    r2 = header_p.add_run("{{STATUS}}")
    _set_run(r2, bold=True, size=12, color=RGBColor(0x2E, 0x8B, 0x57))

    # ---------- SUMMARY SCOPE ----------
    _section_heading(doc, "DIVISION 01 — GENERAL CONDITIONS")
    _item_heading(doc, "00-00 — Summary Scope of Work", "$0.00")
    _para(doc, "Summary Scope of Work", bold=True)
    _para(doc, "{{SUMMARY_SCOPE_OF_WORK}}")
    doc.add_paragraph()
    _para(doc, "Trades involved (preliminary): {{TRADES}}")
    _para(doc, "Bid Due: {{DUE_DATE}}")
    doc.add_paragraph()
    _para(doc, "General")
    _bullets(doc, [
        "All work shall be performed in accordance with industry standards and manufacturer specifications.",
        "Any concealed conditions or required repairs not visible at the time of proposal shall be addressed on a Time and Material basis upon approval.",
    ])

    # ---------- STANDARD GENERAL CONDITIONS BOILERPLATE ----------
    _item_heading(doc, "01-04 — Portable Bathroom")
    _para(doc, "General Specification:")
    _bullets(doc, [
        "Provide and maintain one (1) portable toilet for the duration of the project.",
        "Locate unit in a safe and accessible area as determined by the Contractor.",
        "Maintain unit in clean and sanitary condition at all times.",
        "Provide regular servicing, including cleaning and restocking, on a weekly basis.",
    ])
    _para(doc, "Exclusions & Qualifications:")
    _bullets(doc, [
        "Additional servicing beyond standard weekly maintenance shall be billed as an additional cost.",
        "Relocation of the unit after initial placement, if requested by Owner, may result in additional charges.",
    ])

    _item_heading(doc, "01-05 — Temporary Weather Protection")
    _para(doc, "General Specification:")
    _bullets(doc, [
        "Provide temporary weather protection in all areas of active work to safeguard materials, equipment, and existing conditions.",
        "Install and maintain protection to prevent moisture intrusion, weather damage, and delays due to exposure.",
        "Methods may include roof tarping, floor protection, and temporary enclosures as required.",
    ])
    _para(doc, "Exclusions & Qualifications:")
    _bullets(doc, [
        "Temporary protection is limited to standard construction practices and duration of active work in affected areas.",
        "Extended duration, abnormal weather conditions, or Owner-requested additional protection shall be treated as a Change Order.",
        "Contractor is not responsible for damage resulting from extreme weather events beyond reasonable temporary protection measures.",
    ])

    _item_heading(doc, "01-06 — Weekly and Final Cleaning")
    _para(doc, "General Specification:")
    _bullets(doc, [
        "Maintain a clean and organized job site at all times.",
        "Broom-sweep all active work areas at the close of each week.",
        "Remove packaging, crating materials, and debris daily; no accumulation permitted overnight.",
        "Provide one professional post-construction cleaning upon completion to prepare the home for occupancy.",
    ])
    _para(doc, "Exclusions & Qualifications:")
    _bullets(doc, [
        "Normal post-construction dust migration may continue for 2–3 months after occupancy and is not considered a defect.",
        "Additional or repeat cleanings requested by Owner shall be billed at $180.00 per hour.",
        "Duct cleaning and specialized air quality remediation are excluded unless specifically noted.",
    ])

    _item_heading(doc, "01-07 — Tree Protection", "$0.00")
    _para(doc, "General Specification:")
    _bullets(doc, [
        "Tree protection is not included in this Contract.",
        "Contractor shall make reasonable efforts to avoid unnecessary disturbance to existing trees and landscaped areas within the limits of the Work.",
    ])
    _para(doc, "Exclusions & Qualifications:")
    _bullets(doc, [
        "Installation of tree protection measures, including fencing, barriers, or root protection, is excluded.",
        "Compliance with requirements from any governing authority is excluded unless specifically noted.",
        "If required, Contractor shall provide a separate proposal based on project requirements.",
        "All landscaping, including tree removal, planting, topsoil, seeding, and irrigation, is by Others.",
    ])

    # ---------- PER-PROJECT SCOPE PLACEHOLDER ----------
    doc.add_page_break()
    _section_heading(doc, "DIVISION 02 — 09 — PER-PROJECT SCOPE  (price these line items)")
    _para(doc, "[ Drafted from incoming scope/drawings; review and price ]", bold=True)
    _para(doc, "{{SUMMARY_SCOPE_OF_WORK}}")

    # ---------- EXCLUSIONS & QUALIFICATIONS PART I-VI ----------
    doc.add_page_break()
    _section_heading(doc, "DIVISION 32 — EXCLUSIONS & QUALIFICATIONS")

    _item_heading(doc, "32-00 — Part I — General + Payment", "$0.00")
    _bullets(doc, [
        "This proposal shall be read in conjunction with the applicable AIA Owner-Contractor Agreement for this Project.",
        "Any work not specifically outlined in this proposal is excluded.",
        "Any alterations or deviations involving additional cost shall be treated as Additional Work and executed only upon written approval.",
        "This proposal is based on visual review of accessible conditions only.",
    ])
    _para(doc, "Payment Terms:")
    _bullets(doc, [
        "A deposit of thirty percent (30%) of the Contract Sum is due upon execution of the Agreement.",
        "Progress payments shall be made as invoiced based on Work in place and materials procured.",
        "Invoices are due upon receipt. No retainage shall be withheld.",
        "Final payment is due upon Substantial Completion.",
        "Contractor reserves the right to suspend Work in the event payments are not made in accordance with these terms.",
        "Owner shall be responsible for all costs of collection, including reasonable attorneys' fees, interest, and administrative costs.",
    ])

    _item_heading(doc, "32-00 — Part II — Allowances / Escalation / Delays", "$0.00")
    _para(doc, "Allowances:")
    _bullets(doc, [
        "Allowances are estimated amounts only and shall be reconciled based on actual cost of labor, materials, and installation.",
        "Allowance overages shall constitute Additional Work and require approval prior to proceeding.",
        "Unused Allowance amounts shall be credited to Owner.",
    ])
    _para(doc, "Material Escalation / Tariffs:")
    _bullets(doc, [
        "This proposal is based on current material pricing.",
        "Contractor is not responsible for increases due to market conditions, tariffs, supply chain disruption, or manufacturer pricing changes.",
    ])
    _para(doc, "Owner-Caused Delays / Direct-Pay Items:")
    _bullets(doc, [
        "Delays caused by Owner decisions, selections, approvals, or Owner-supplied materials shall result in additional time and cost.",
        "Contractor assumes no responsibility for work, delays, or defects caused by Owner-retained vendors.",
    ])

    _item_heading(doc, "32-00 — Part III — Existing Conditions / Permits / Demo", "$0.00")
    _bullets(doc, [
        "Contractor is not responsible for concealed conditions including rot, mold, structural deficiencies, insulation issues, or hidden damage.",
        "Any rot discovered shall be repaired on a Time and Material basis upon approval.",
        "Permits are not included and are the responsibility of the Owner.",
        "No Architectural or Engineering services are included.",
        "Owner shall remove all personal property from work areas prior to commencement.",
    ])

    _item_heading(doc, "32-00 — Part IV — Roofing / Windows / Framing", "$0.00")
    _bullets(doc, [
        "Scope is limited to Work specifically described in this proposal.",
        "Flat roofs, balconies, and non-scoped areas are excluded unless noted.",
        "Roof decking or structural repairs are excluded unless noted.",
        "Contractor is not responsible for removal or reinstallation of solar, satellite, or similar equipment.",
        "Windows furnished by Others carry no Contractor warranty for unit, glass, or hardware defects.",
        "Return trips due to Owner-supplied materials shall be billed as Additional Work.",
    ])

    _item_heading(doc, "32-00 — Part V — Trim / Siding / Interior / Masonry", "$0.00")
    _bullets(doc, [
        "Allowance work is limited to areas impacted by the Work.",
        "Full replacement of trim, siding, or veneer systems is excluded.",
        "Exact match of materials, color, profile, and finish is not guaranteed.",
        "Substrate or structural repairs beyond minor work are excluded.",
        "Painting is excluded unless specifically noted.",
    ])

    _item_heading(doc, "32-00 — Part VI — Paint / Environmental / General", "$0.00")
    _bullets(doc, [
        "Painting is limited to touch-ups in areas affected by the Work only.",
        "Contractor does not test for or remediate asbestos, lead, mold, or hazardous materials.",
        "Mold testing, if included, is limited to sampling only; remediation is excluded.",
        "Work is subject to weather, material availability, and conditions beyond Contractor control.",
    ])

    # ---------- DIVISION 33 LIMITATION OF LIABILITY ----------
    _section_heading(doc, "DIVISION 33 — LIMITATION OF LIABILITY")
    _bullets(doc, [
        "Contractor's total liability for any claims, damages, losses, or expenses arising out of or relating to the Work shall be limited to the total Contract Sum.",
        "Contractor shall not be liable for indirect, incidental, consequential, or special damages, including loss of use, income, profits, or diminution in property value.",
        "Contractor shall not be responsible for pre-existing conditions, concealed conditions, or latent defects within the structure.",
        "Contractor shall have no liability for materials, equipment, or systems supplied by the Owner or installed by others.",
        "Contractor does not guarantee exact matching of existing materials, finishes, colors, or textures.",
        "Contractor shall not be responsible for environmental conditions including mold, mildew, air quality, allergens, or moisture-related issues.",
        "Contractor shall not be liable for damage caused by weather events during active construction phases.",
        "Contractor warranty applies only to labor performed under this contract. No warranty is provided for Owner-supplied materials, work by others, pre-existing conditions, or normal material movement.",
    ])

    # ---------- DIVISION 34 SCOPE DEFINITION ----------
    _section_heading(doc, "DIVISION 34 — SCOPE DEFINITION & OWNERSHIP")
    _bullets(doc, [
        "The Scope of Work is limited strictly to the items specifically described in this proposal.",
        "Any work, material, detail, or condition not explicitly stated shall be excluded.",
        "In the event of ambiguity, omission, or conflict within the scope, Contractor's interpretation shall govern unless otherwise agreed in writing.",
        "Contractor shall have sole control over construction means, methods, techniques, sequences, and procedures.",
        "Contractor is not responsible for coordination, performance, or scheduling of work performed by others.",
        "Any portion of the Work identified as an Allowance shall be adjusted based on actual conditions and requirements.",
        "All exclusions listed within this proposal are incorporated into and govern the Scope of Work.",
    ])

    # ---------- PAYMENT SCHEDULE ----------
    doc.add_page_break()
    _section_heading(doc, "PAYMENT SCHEDULE")
    pay = doc.add_table(rows=3, cols=2)
    pay.rows[0].cells[0].paragraphs[0].add_run("Name").bold = True
    pay.rows[0].cells[1].paragraphs[0].add_run("Amount").bold = True
    pay.rows[1].cells[0].text = "Deposit (30%)"
    pay.rows[1].cells[1].text = "$ TBD"
    pay.rows[2].cells[0].text = "Progress / Final"
    pay.rows[2].cells[1].text = "$ TBD"

    # ---------- SIGNATURE BLOCK ----------
    doc.add_paragraph()
    _para(doc, "The above specifications, costs, and terms are hereby accepted.", bold=True)
    doc.add_paragraph()
    doc.add_paragraph("__________________________________________            ____________________")
    _para(doc, "{{CONTACT_NAME}}                                                                            DATE")
    doc.add_paragraph()
    doc.add_paragraph("__________________________________________            ____________________")
    _para(doc, "{{SENDER_NAME}} — {{SENDER_COMPANY}}                                            DATE")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
