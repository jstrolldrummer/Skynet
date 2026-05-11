from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path
from typing import Mapping, Optional

from docx import Document

from bid_watcher.config import Settings
from bid_watcher.models import BidClassification, IncomingEmail
from bid_watcher.util.logging import get_logger

log = get_logger(__name__)

# Placeholder syntax in the template: {{CLIENT_COMPANY}}, {{PROJECT_NAME}}, etc.
PLACEHOLDER = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}")


def _format_long_date(d: date) -> str:
    return f"{d.strftime('%B')} {d.day}, {d.year}"


def _replace_in_paragraph(paragraph, mapping: Mapping[str, str]) -> None:
    if "{{" not in paragraph.text:
        return
    # Collapse runs so {{ }} that span runs still get replaced
    full = paragraph.text
    new = PLACEHOLDER.sub(lambda m: mapping.get(m.group(1), m.group(0)), full)
    if new == full:
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new
    else:
        paragraph.add_run(new)


def _fill_doc(doc: Document, mapping: Mapping[str, str]) -> None:
    for p in doc.paragraphs:
        _replace_in_paragraph(p, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, mapping)


def build_placeholder_map(
    email: IncomingEmail,
    c: BidClassification,
    settings: Optional[Settings] = None,
) -> dict[str, str]:
    project_title = c.project_name or email.subject or "New Project"
    if c.project_location:
        project_title = f"{project_title} — {c.project_location}"
    return {
        # Sender / company (cover page)
        "SENDER_NAME": getattr(settings, "sender_name", "") if settings else "",
        "SENDER_COMPANY": getattr(settings, "sender_company", "") if settings else "",
        "SENDER_PHONE": getattr(settings, "sender_phone", "") if settings else "",
        "SENDER_EMAIL": getattr(settings, "sender_email", "") if settings else "",
        "SENDER_ADDRESS": getattr(settings, "sender_address", "") if settings else "",
        # Client / project
        "CLIENT_COMPANY": c.client_company or "",
        "CONTACT_NAME": c.contact_name or email.sender_name or "",
        "CONTACT_EMAIL": email.sender_email,
        "PROJECT_NAME": c.project_name or "",
        "PROJECT_LOCATION": c.project_location or "",
        "PROJECT_TITLE": project_title,
        "PROPOSAL_NUMBER": "Proposal 1-1 (DRAFT)",
        "STATUS": "DRAFT",
        "DUE_DATE": c.due_date or "",
        "SCOPE_SUMMARY": c.scope_summary or "",
        "SUMMARY_SCOPE_OF_WORK": c.scope_summary or "(scope to be priced)",
        "TRADES": ", ".join(c.trades),
        "TODAY": date.today().isoformat(),
        "ISSUE_DATE": _format_long_date(date.today()),
        "RECEIVED_DATE": email.received_at.date().isoformat(),
        "SUBJECT": email.subject,
    }


def render_proposal(
    template_path: Path,
    email: IncomingEmail,
    c: BidClassification,
    settings: Optional[Settings] = None,
) -> bytes:
    if not template_path.exists():
        log.warning("proposal.template_missing", path=str(template_path))
        return _render_fallback(email, c, settings)
    doc = Document(str(template_path))
    _fill_doc(doc, build_placeholder_map(email, c, settings))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_fallback(email: IncomingEmail, c: BidClassification, settings: Optional[Settings] = None) -> bytes:
    """Minimal proposal when no template is configured yet."""
    doc = Document()
    doc.add_heading("PROPOSAL — DRAFT", level=0)
    mapping = build_placeholder_map(email, c, settings)
    for label, key in [
        ("Client", "CLIENT_COMPANY"),
        ("Contact", "CONTACT_NAME"),
        ("Project", "PROJECT_NAME"),
        ("Location", "PROJECT_LOCATION"),
        ("Due", "DUE_DATE"),
        ("Date", "TODAY"),
    ]:
        doc.add_paragraph(f"{label}: {mapping[key]}")
    doc.add_heading("Scope Summary", level=1)
    doc.add_paragraph(c.scope_summary or "(scope to be filled in)")
    doc.add_heading("Trades", level=1)
    doc.add_paragraph(", ".join(c.trades) if c.trades else "(to be determined)")
    doc.add_heading("Notes from email", level=1)
    doc.add_paragraph(email.subject)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
